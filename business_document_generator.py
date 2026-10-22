"""
Business Document Generator
Generates Word (.docx) documents from text and structured inputs.
"""

import json
import os

from docx import Document
from docx.shared import Pt


def generate_word_document(
    title: str, paragraphs: list, output_filename: str = "generated_doc.docx"
) -> str:
    """
    Creates a professionally formatted Microsoft Word document.

    Args:
        title (str): The main title of the document.
        paragraphs (list): A list of strings, each representing a paragraph.
        output_filename (str): The path to save the .docx file.

    Returns:
        JSON string containing the absolute path to the generated file or an error.
    """
    try:
        doc = Document()

        # Add Title
        title_run = doc.add_heading(title, 0).runs[0]
        title_run.font.size = Pt(24)

        # Add Paragraphs
        for p_text in paragraphs:
            p = doc.add_paragraph(p_text)
            p.style.font.size = Pt(12)

        # Save Document
        abs_path = os.path.abspath(output_filename)
        doc.save(abs_path)

        return json.dumps(
            {
                "status": "success",
                "message": "Document generated successfully.",
                "file_path": abs_path,
            },
            indent=2,
        )

    except Exception as e:
        return json.dumps({"status": "error", "message": str(e)})


if __name__ == "__main__":
    import sys
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

    test_title = "Non-Disclosure Agreement (NDA)"
    test_paras = [
        "This Non-Disclosure Agreement is entered into by and between the parties.",
        "Confidential Information includes all data, materials, products, technology, computer programs, specifications, manuals, business plans, software, marketing plans, financial information, and other information disclosed or submitted, orally, in writing, or by any other media.",
        "Signature: ___________________________ Date: _______________",
    ]
    print(generate_word_document(test_title, test_paras, "test_nda.docx"))
